from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Report" / "中证A500指数增强基金研究报告底稿.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=110, start=110, bottom=110, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, latin="Arial Unicode MS", east="Arial Unicode MS", size=10.5, bold=False, color="000000"):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(6)
    if bold_lead and text.startswith(bold_lead):
        a = p.add_run(bold_lead)
        set_font(a, bold=True)
        b = p.add_run(text[len(bold_lead):])
        set_font(b)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, size=16 if level == 1 else 13, bold=True)
    return p


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.4)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.6)
sec.right_margin = Cm(2.4)

styles = doc.styles
styles["Normal"].font.name = "Arial Unicode MS"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
styles["Normal"].font.size = Pt(10.5)
for name in ["Title", "Heading 1", "Heading 2"]:
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(105)
title.paragraph_format.space_after = Pt(22)
r = title.add_run("中证 A500 指数增强基金设计与实证研究")
set_font(r, size=24, bold=True)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("第一次小组作业报告底稿")
set_font(r, size=14)

meta = doc.add_table(rows=5, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
labels = ["课程", "小组", "成员", "提交日期", "正式回测状态"]
values = ["硕士组合管理", "第    组", "填写姓名与学号", "2026 年    月    日", "导入 Wind 数据后更新"]
for i, (label, value) in enumerate(zip(labels, values)):
    meta.rows[i].cells[0].width = Cm(4)
    meta.rows[i].cells[1].width = Cm(9)
    for cell in meta.rows[i].cells:
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(meta.rows[i].cells[0], "D9EAF2")
    p = meta.rows[i].cells[0].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(label), bold=True)
    p = meta.rows[i].cells[1].paragraphs[0]
    set_font(p.add_run(value))

doc.add_page_break()
add_heading(doc, "摘要", 1)
add_body(doc, "本文以中证 A500 指数为基准，构建周频调仓的股票型指数增强组合。策略使用价值、质量、动量和低波动因子形成主动收益观点，通过 Black-Litterman 模型与市场均衡收益融合，再使用带跟踪误差、行业偏离、个股偏离和换手约束的 Markowitz 模型确定组合权重。正式 Wind 数据导入后，本节将自动汇总样本期、年化收益、超额收益、信息比率及实现跟踪误差。")
add_body(doc, "关键词：中证 A500；指数增强；多因子模型；Black-Litterman；Markowitz；跟踪误差")

add_heading(doc, "一 我国股票型指数增强基金概况", 1)
add_body(doc, "指数增强基金在控制相对基准风险的前提下，通过股票选择、行业配置、交易执行或衍生工具争取超额收益。报告应重点说明国内股票型指数增强基金的产品形态、收益来源、风险控制和评价方法，并使用中证指数公司、基金合同、招募说明书和定期报告等原始资料核验事实。")
add_heading(doc, "中证 A500 指数及相关产品", 2)
add_body(doc, "本节填写指数编制目标、样本空间、行业覆盖、权重规则和调整机制，并比较具有可比性的中证 A500 指数增强产品。产品规模、费率和业绩信息必须注明统计日期，避免混用不同份额或不同收益口径。")

add_heading(doc, "二 数据与处理", 1)
add_body(doc, "研究数据从 Wind 终端手工导出。原始文件不做数值修改，清洗程序从 CSV 或 XLSX 开始运行。所有信号仅使用调仓日前已经公开的数据，财务指标自公告日后的下一交易日起生效。")
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
headers = ["数据表", "主要内容", "频率", "防前视处理"]
for j, text in enumerate(headers):
    set_cell_shading(table.rows[0].cells[j], "1F4E79")
    p = table.rows[0].cells[j].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), size=9.5, bold=True, color="FFFFFF")
set_repeat_table_header(table.rows[0])
rows = [
    ("benchmark", "价格指数与全收益指数", "日频", "使用当日已公布收盘值"),
    ("constituents", "历史成分股与权重", "周频", "只用最近生效的历史快照"),
    ("prices", "前复权行情与交易状态", "日频", "信号日收盘后计算"),
    ("industry", "申万一级行业", "事件频率", "按生效日向后匹配"),
    ("valuation", "PE PB PS 股息率", "周频", "只取信号日前记录"),
    ("fundamentals", "质量和杠杆指标", "报告期", "按公告日匹配并限制陈旧期"),
]
for i, row in enumerate(rows):
    cells = table.add_row().cells
    for j, value in enumerate(row):
        set_cell_margins(cells[j], top=90, bottom=90)
        if i % 2:
            set_cell_shading(cells[j], "F3F7FA")
        p = cells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(value), size=9)

add_heading(doc, "三 模型设计", 1)
add_heading(doc, "多因子预期收益", 2)
add_body(doc, "价值因子由 EP、BP、SP 和股息率构成；质量因子由 ROE、毛利率、经营现金流质量和低杠杆构成；动量使用过去 12 个月收益并剔除最近 1 个月；低波动因子使用近 60 个交易日波动率和下行波动率。每周截面进行 1% 和 99% 去极值及行业内标准化，滚动 Rank IC 和 ICIR 决定因子权重。")
add_heading(doc, "Black-Litterman 后验收益", 2)
add_body(doc, "模型以基准权重和协方差矩阵推导市场均衡收益，将综合因子得分映射为逐股相对收益观点，并依据滚动 IC 稳定性设置置信度。后验收益同时保留基准隐含信息与因子信号，供组合优化使用。")
add_heading(doc, "带约束的 Markowitz 优化", 2)
add_body(doc, "目标函数最大化预期主动收益，并惩罚主动风险与换手。组合仅允许多头且保持满仓；年化预期跟踪误差不超过 6%；申万一级行业主动偏离不超过 2 个百分点；个股主动权重偏离不超过 1 个百分点；单周单边换手率不超过 20%。优化器每期保存状态并重新检查硬约束，失败时终止回测。")

add_heading(doc, "四 回测设计", 1)
add_body(doc, "每周最后一个交易日收盘后计算信号，下一个交易日执行。主结果采用单边 20bp 交易成本，并分别测试 10bp 和 30bp。停牌或涨跌停股票保持原权重，其余股票在剩余可行空间内优化。协方差矩阵使用近 120 个交易日收益估计。")

add_heading(doc, "五 实证结果", 1)
add_body(doc, "正式 Wind 回测完成后，从 Result/official 导入绩效表和图表。本节必须报告累计收益、年化收益、年化超额收益、波动率、Sharpe、最大回撤、信息比率、实现跟踪误差、胜率、换手率和交易成本。")
result_table = doc.add_table(rows=1, cols=4)
result_table.style = "Table Grid"
for j, h in enumerate(["指标", "完整策略", "仅多因子", "中证 A500"]):
    set_cell_shading(result_table.rows[0].cells[j], "1F4E79")
    p = result_table.rows[0].cells[j].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(h), size=9.5, bold=True, color="FFFFFF")
for metric in ["年化收益", "年化超额收益", "年化波动率", "最大回撤", "信息比率", "实现跟踪误差", "平均单边换手率"]:
    cells = result_table.add_row().cells
    for j, value in enumerate([metric, "正式回测后填入", "正式回测后填入", "正式回测后填入"]):
        set_cell_margins(cells[j], top=90, bottom=90)
        p = cells[j].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(value), size=9)

add_heading(doc, "六 稳健性检验", 1)
add_body(doc, "比较交易成本 10bp、20bp 和 30bp，跟踪误差预算 4%、5% 和 6%，逐类剔除因子，并增加不使用 Black-Litterman 的对照组合。报告全部预设实验，不依据结果选择性删除。")

add_heading(doc, "七 结论与局限", 1)
add_body(doc, "结论区分样本内分析与样本外表现，并讨论中证 A500 历史长度、财务数据更新频率、交易成本假设、协方差估计误差和优化参数敏感性。正式结论应由实证结果支持。")

add_heading(doc, "八 小组成员分工", 1)
roles = doc.add_table(rows=1, cols=4)
roles.style = "Table Grid"
for j, h in enumerate(["成员", "负责内容", "具体文件或图表", "交叉核验"]):
    set_cell_shading(roles.rows[0].cells[j], "1F4E79")
    p = roles.rows[0].cells[j].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(h), size=9.5, bold=True, color="FFFFFF")
for label, work in [("A", "产品与文献研究"), ("B", "Wind 行情与成分"), ("C", "财务数据与清洗"), ("D", "因子与 IC 检验"), ("E", "回测与绩效评价"), ("F", "BL 模型与优化")]:
    cells = roles.add_row().cells
    for j, value in enumerate([label, work, "填写实际产出", "填写核验对象"]):
        set_cell_margins(cells[j], top=90, bottom=90)
        p = cells[j].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(value), size=9)

add_heading(doc, "九 AI 工具使用说明", 1)
add_body(doc, "本项目使用 OpenAI Codex 协助设计研究流程、搭建代码框架、生成文档和演示模板，并辅助检查代码结构。小组成员负责从 Wind 导出数据、核验字段与来源、运行和检查全部代码、解释实证结果、补充文献及完成最终写作。报告中的数据、结论和引用均由小组成员从原始来源复核。")

add_heading(doc, "参考文献", 1)
add_body(doc, "按照课程要求统一格式列出中证指数公司资料、基金文件、论文、券商研报和网站资料。正文引用与文末条目一一对应，并注明网页访问日期和数据统计日期。")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("中证 A500 指数增强基金研究报告"), size=8, color="666666")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
